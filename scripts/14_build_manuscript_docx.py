#!/usr/bin/env python3
"""Build a review-ready Word manuscript from the Markdown draft."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT = Path(__file__).resolve().parents[1]
SOURCE = PROJECT / "manuscript" / "MANUSCRIPT_DRAFT_V1.md"
OUTPUT = PROJECT / "manuscript" / "MANUSCRIPT_DRAFT_V1.docx"
FIGURE_DIR = PROJECT / "figures" / "publication"

FIGURES = {
    1: FIGURE_DIR / "Figure_1_study_design_and_QC.png",
    2: FIGURE_DIR / "Figure_2_Cd33_cross_dose_effects.png",
    3: FIGURE_DIR / "Figure_3_gene_level_meta_analysis.png",
    4: FIGURE_DIR / "Figure_4_pathway_reversal_test.png",
    5: FIGURE_DIR / "Figure_5_human_external_validation.png",
}
SUPP_FIGURE_DIR = PROJECT / "figures" / "robust_gene_enrichment"

FONT = "Times New Roman"
BLUE = "1F4E78"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F4F4F4"
BORDER = "D9D9D9"


def set_run_font(run, size: float | None = None, bold=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + edge
        element = tc_mar.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn("w:" + edge))
        if element is None:
            element = OxmlElement("w:" + edge)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_run_font(run, 10)


def add_hyperlink(paragraph, text: str, url: str, bold=False, italic=False, superscript=False):
    part = paragraph.part
    relation_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia"):
        r_fonts.set(qn("w:" + attr), FONT)
    r_pr.append(r_fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    if bold:
        r_pr.append(OxmlElement("w:b"))
    if italic:
        r_pr.append(OxmlElement("w:i"))
    if superscript:
        vert = OxmlElement("w:vertAlign")
        vert.set(qn("w:val"), "superscript")
        r_pr.append(vert)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "24")
    r_pr.append(size)
    run.append(r_pr)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


TOKEN_PATTERN = re.compile(
    r"(<sup>.*?</sup>|\[[^\]]+\]\([^\)]+\)|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)"
)


def add_inline(paragraph, text: str, size: float = 12, base_bold=False, base_italic=False):
    """Add basic Markdown inline formatting without altering visible text."""
    pos = 0
    for match in TOKEN_PATTERN.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_run_font(run, size, base_bold, base_italic)
        token = match.group(0)
        if token.startswith("<sup>"):
            visible = token[5:-6]
            run = paragraph.add_run(visible)
            set_run_font(run, max(size - 2, 8), base_bold, base_italic)
            run.font.superscript = True
        elif token.startswith("["):
            link_match = re.fullmatch(r"\[([^\]]+)\]\(([^\)]+)\)", token)
            if link_match:
                add_hyperlink(
                    paragraph,
                    link_match.group(1),
                    link_match.group(2),
                    base_bold,
                    base_italic,
                )
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size, True, base_italic)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size, base_bold, base_italic)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size, base_bold, True)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size, base_bold, base_italic)


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)
    pf.widow_control = True

    title = styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title.font.size = Pt(16)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.line_spacing = 1.15
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 14, 12, 6),
        ("Heading 2", 12, 10, 4),
        ("Heading 3", 11, 8, 3),
    ):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.15


def configure_section(section, landscape=False, table_section=False):
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        margin = Inches(0.65 if table_section else 0.75)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        margin = Inches(1)
    section.top_margin = margin
    # Leave extra clearance above the centred footer page number.  LibreOffice's
    # pagination is slightly less conservative than Word's on text-dense pages,
    # so this margin prevents the final body line from approaching the footer.
    section.bottom_margin = Inches(1.55) if not landscape else margin
    section.left_margin = margin
    section.right_margin = margin
    section.header_distance = Inches(0.35)
    # Keep the page number low in the footer so it remains visually distinct
    # from a full final body line on densely paginated pages.
    section.footer_distance = Inches(0.05)


def ensure_page_number(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.clear()
    add_page_field(paragraph)


def clean_cell_text(value: str) -> str:
    return value.strip().replace("<br>", "\n")


def prevent_row_split(row):
    """Keep a table row together so continuation pages do not show cell fragments."""
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def add_table(doc: Document, rows: list[list[str]], table_number: int | None):
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    available = 9.7
    ratios = {
        1: [1.15, 1.05, 1.15, 0.70, 0.80, 1.20, 1.10, 0.70],
        2: [1.25, 1.30, 0.80, 1.00, 1.15, 4.20],
        3: [1.20, 1.65, 1.80, 1.35, 0.75, 0.75, 2.20],
    }.get(table_number, [1.0] * columns)
    if len(ratios) != columns:
        ratios = [1.0] * columns
    widths = [available * ratio / sum(ratios) for ratio in ratios]

    for row_idx, values in enumerate(rows):
        prevent_row_split(table.rows[row_idx])
        for col_idx in range(columns):
            cell = table.rows[row_idx].cells[col_idx]
            cell.width = Inches(widths[col_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
                if col_idx in (0, 1, columns - 1)
                else WD_ALIGN_PARAGRAPH.CENTER
            )
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_before = Pt(1.5)
            paragraph.paragraph_format.space_after = Pt(1.5)
            value = clean_cell_text(values[col_idx]) if col_idx < len(values) else ""
            add_inline(paragraph, value, size=8.5, base_bold=(row_idx == 0))
            if row_idx == 0:
                set_cell_shading(cell, BLUE)
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
            elif row_idx % 2 == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            else:
                set_cell_shading(cell, "FFFFFF")
    repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_figure(doc: Document, number: int):
    path = FIGURES[number]
    if not path.exists():
        raise FileNotFoundError(path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(6.35))


def add_body_paragraph(
    doc: Document,
    text: str,
    section_name: str,
    front_matter=False,
    table_note=False,
):
    if section_name == "References" and re.match(r"^\d+\.\s", text):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.keep_together = True
        add_inline(paragraph, text, size=10.5)
        return paragraph

    list_match = re.match(r"^(\d+)\.\s+(.*)$", text)
    bullet_match = re.match(r"^-\s+(.*)$", text)
    if list_match:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.28)
        paragraph.paragraph_format.first_line_indent = Inches(-0.28)
        add_inline(paragraph, f"{list_match.group(1)}. {list_match.group(2)}", size=12)
    elif bullet_match:
        paragraph = doc.add_paragraph(style="List Bullet")
        add_inline(paragraph, bullet_match.group(1), size=12)
    else:
        paragraph = doc.add_paragraph()
        add_inline(paragraph, text, size=11 if front_matter else 12)

    # Body paragraphs must be allowed to break across pages.  Keeping every
    # paragraph together can make LibreOffice overflow a long paragraph into
    # the footer instead of paginating it normally.
    paragraph.paragraph_format.keep_together = False

    if table_note:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.keep_together = True
        for run in paragraph.runs:
            run.font.size = Pt(9.5)
    elif front_matter:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(3)
    elif section_name == "Figure legends":
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(6)
        for run in paragraph.runs:
            run.font.size = Pt(10.5)
    elif text.startswith("[") and "To be completed" in text:
        paragraph.paragraph_format.line_spacing = 1.15
        for run in paragraph.runs:
            run.italic = True
            run.font.color.rgb = RGBColor(120, 60, 60)
    return paragraph


def build_document():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    doc.core_properties.title = "Transcriptomic effects of Cd33-targeting ASO in APP/PS1 mouse brain"
    doc.core_properties.subject = "SCI manuscript draft"
    doc.core_properties.author = "Heng Xu; Huanhuan Ma; Xiaodan Wu"
    doc.core_properties.comments = "Generated from MANUSCRIPT_DRAFT_V1.md"

    configure_styles(doc)
    configure_section(doc.sections[0], landscape=False)
    ensure_page_number(doc.sections[0])

    current_section = ""
    front_matter = True
    table_number = None
    in_landscape = False
    pending_figure = None
    pending_supp_figure = None
    table_rows: list[list[str]] = []
    table_note_pending = False

    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line or line == "---":
            i += 1
            continue

        if line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            add_inline(paragraph, line[2:].strip(), size=16, base_bold=True)
            i += 1
            continue

        if line.startswith("## "):
            heading = line[3:].strip()
            current_section = re.sub(r"^\d+\.\s*", "", heading)
            front_matter = False

            table_match = re.match(r"Table\s+(\d+)\.", heading)
            if table_match and not in_landscape:
                section = doc.add_section(WD_SECTION.NEW_PAGE)
                configure_section(section, landscape=True, table_section=True)
                ensure_page_number(section)
                in_landscape = True
            if heading == "Figure legends" and in_landscape:
                section = doc.add_section(WD_SECTION.NEW_PAGE)
                configure_section(section, landscape=False)
                ensure_page_number(section)
                in_landscape = False

            style = "Heading 1"
            paragraph = doc.add_paragraph(style=style)
            add_inline(paragraph, heading, size=14, base_bold=True)
            if heading == "5. Conclusions":
                paragraph.paragraph_format.page_break_before = True
            table_number = int(table_match.group(1)) if table_match else None
            i += 1
            continue

        if line.startswith("### "):
            heading = line[4:].strip()
            paragraph = doc.add_paragraph(style="Heading 2")
            add_inline(paragraph, heading, size=12, base_bold=True)
            figure_match = re.match(r"(?:Figure|\*\*Fig\.)\s*(\d+)[\*\.]?", heading)
            supp_figure_match = re.match(r"(?:Supplementary Figure (?:S1|S\d+)|\*\*Supplementary Fig\. S1\*\*)", heading)
            pending_figure = int(figure_match.group(1)) if figure_match else None
            pending_supp_figure = True if supp_figure_match else None
            if pending_figure is not None and pending_figure > 1:
                paragraph.paragraph_format.page_break_before = True
            i += 1
            continue

        if line.startswith("#### "):
            paragraph = doc.add_paragraph(style="Heading 3")
            add_inline(paragraph, line[5:].strip(), size=11, base_bold=True)
            i += 1
            continue

        if line.startswith("|"):
            table_rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_line = lines[i].strip()
                values = [cell.strip() for cell in table_line.strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in values):
                    table_rows.append(values)
                i += 1
            add_table(doc, table_rows, table_number)
            table_note_pending = True
            continue

        paragraph = add_body_paragraph(
            doc,
            line,
            current_section,
            front_matter,
            table_note=table_note_pending,
        )
        table_note_pending = False
        if pending_figure is not None:
            add_figure(doc, pending_figure)
            paragraph.paragraph_format.keep_with_next = True
            pending_figure = None
        if pending_supp_figure is not None:
            supp_path = SUPP_FIGURE_DIR / "top_CD33_robust_GO_BP.png"
            p2 = doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_before = Pt(6)
            p2.paragraph_format.space_after = Pt(6)
            run = p2.add_run()
            run.add_picture(str(supp_path), width=Inches(6.35))
            pending_supp_figure = None
        i += 1

    for section in doc.sections:
        section.start_type = WD_SECTION.NEW_PAGE if section is not doc.sections[0] else section.start_type

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
